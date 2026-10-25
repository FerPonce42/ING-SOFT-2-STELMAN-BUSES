import unittest
from io import BytesIO
from unittest.mock import Mock

from openpyxl import load_workbook

from app.services.reporte_service import ReporteService


def _esquema(tablas):
    """Esquema falso. tablas: {nombre: {"ddl", "columnas", "fks", "total", "filas"}}"""
    esquema = Mock()
    esquema.nombre_bd.return_value = "stelman_db"
    esquema.tablas.return_value = list(tablas)
    esquema.ddl.side_effect = lambda t: tablas[t]["ddl"]
    esquema.columnas.side_effect = lambda t: tablas[t]["columnas"]
    esquema.claves_foraneas.side_effect = lambda t: tablas[t].get("fks", [])
    esquema.contar.side_effect = lambda t: tablas[t].get("total", len(tablas[t].get("filas", [])))
    esquema.filas.side_effect = lambda t: tablas[t].get("filas", [])
    return esquema


def _col(nombre, clave=""):
    return {"COLUMN_NAME": nombre, "COLUMN_KEY": clave}


def _fk(columna, tabla, destino):
    return {"COLUMN_NAME": columna, "REFERENCED_TABLE_NAME": tabla, "REFERENCED_COLUMN_NAME": destino}


TABLAS = {
    "ruta": {
        "ddl": "CREATE TABLE `ruta` (id int)",
        "columnas": [_col("id", "PRI"), _col("nombre")],
        "filas": [{"id": 1, "nombre": "Arequipa-Juliaca"}],
    },
    "bus": {
        "ddl": "CREATE TABLE `bus` (id int, id_ruta int)",
        "columnas": [_col("id", "PRI"), _col("id_ruta", "MUL")],
        "fks": [_fk("id_ruta", "ruta", "id")],
        "filas": [{"id": 10, "id_ruta": 1}, {"id": 11, "id_ruta": 1}],
    },
}


class DocumentacionExcelTest(unittest.TestCase):
    def setUp(self):
        self.esquema = _esquema(TABLAS)
        self.servicio = ReporteService(self.esquema)
        self.wb = load_workbook(self.servicio.documentacion_excel())

    def tearDown(self):
        self.esquema = self.servicio = self.wb = None

    def test_devuelve_un_buffer_listo_para_leer_desde_el_inicio(self):
        salida = self.servicio.documentacion_excel()
        self.assertIsInstance(salida, BytesIO)
        self.assertEqual(salida.tell(), 0)

    def test_el_contenido_es_un_xlsx_valido(self):
        salida = self.servicio.documentacion_excel()
        self.assertEqual(salida.read(2), b"PK")

    def test_crea_las_hojas_en_el_orden_esperado(self):
        self.assertEqual(
            self.wb.sheetnames,
            ["PORTADA", "Índice", "CREATE", "ruta", "bus", "ERD_TEXTUAL"],
        )

    def test_portada_muestra_el_nombre_de_la_base_de_datos(self):
        self.assertEqual(self.wb["PORTADA"]["A3"].value, "Base de datos: stelman_db")

    def test_portada_tiene_titulo_y_generador(self):
        portada = self.wb["PORTADA"]
        self.assertEqual(portada["A1"].value, "DOCUMENTACIÓN TÉCNICA - BD STELMAN BUSES")
        self.assertEqual(portada["A5"].value, "Generado por: Sistema Stelman")

    def test_indice_tiene_los_encabezados_esperados(self):
        fila = [c.value for c in self.wb["Índice"][1]]
        self.assertEqual(
            fila, ["#", "Tabla", "Atributos", "Registros", "PK", "FK", "Estado", "Link hoja"]
        )

    def test_indice_resume_cada_tabla(self):
        fila = [c.value for c in self.wb["Índice"][3]]
        self.assertEqual(fila[:7], [2, "bus", 2, 2, "id", "id_ruta → ruta.id", "OK"])

    def test_indice_enlaza_a_la_hoja_de_cada_tabla(self):
        self.assertEqual(self.wb["Índice"]["H2"].value, '=HYPERLINK("#ruta!A1","Abrir")')

    def test_indice_usa_el_total_de_registros_del_esquema(self):
        tablas = {"ruta": dict(TABLAS["ruta"], total=5000)}
        wb = load_workbook(ReporteService(_esquema(tablas)).documentacion_excel())
        self.assertEqual(wb["Índice"]["D2"].value, 5000)

    def test_hoja_create_lista_el_ddl_de_cada_tabla(self):
        hoja = self.wb["CREATE"]
        self.assertEqual([c.value for c in hoja[2]], ["ruta", "CREATE TABLE `ruta` (id int)"])
        self.assertEqual(hoja["A3"].value, "bus")

    def test_hoja_de_tabla_muestra_pk_y_fk(self):
        hoja = self.wb["bus"]
        self.assertEqual(hoja["A1"].value, "CREATE TABLE: bus")
        self.assertEqual(hoja["B6"].value, "id")
        self.assertEqual(hoja["B7"].value, "id_ruta → ruta.id")

    def test_hoja_de_tabla_escribe_encabezados_y_datos(self):
        hoja = self.wb["bus"]
        self.assertEqual([hoja.cell(row=9, column=c).value for c in (1, 2)], ["id", "id_ruta"])
        self.assertEqual([hoja.cell(row=10, column=c).value for c in (1, 2)], [10, 1])
        self.assertEqual([hoja.cell(row=11, column=c).value for c in (1, 2)], [11, 1])

    def test_pk_se_resalta_en_amarillo_y_fk_en_azul(self):
        hoja = self.wb["bus"]
        self.assertTrue(hoja["A10"].fill.fgColor.rgb.endswith("FFD966"))
        self.assertTrue(hoja["B10"].fill.fgColor.rgb.endswith("C9DAF8"))

    def test_con_filas_activa_filtro_y_congela_encabezado(self):
        hoja = self.wb["bus"]
        self.assertEqual(hoja.auto_filter.ref, "A9:B11")
        self.assertEqual(hoja.freeze_panes, "A10")

    def test_tabla_sin_filas_no_activa_filtro(self):
        tablas = {"ruta": dict(TABLAS["ruta"], filas=[])}
        wb = load_workbook(ReporteService(_esquema(tablas)).documentacion_excel())
        self.assertIsNone(wb["ruta"].auto_filter.ref)

    def test_nombre_de_hoja_se_recorta_a_31_caracteres(self):
        nombre = "t" * 40
        tablas = {nombre: dict(TABLAS["ruta"])}
        wb = load_workbook(ReporteService(_esquema(tablas)).documentacion_excel())
        self.assertIn("t" * 31, wb.sheetnames)
        self.assertNotIn(nombre, wb.sheetnames)

    def test_erd_textual_lista_las_relaciones(self):
        self.assertEqual(self.wb["ERD_TEXTUAL"]["A3"].value, "bus.id_ruta  →  ruta.id")

    def test_sin_relaciones_el_erd_lo_indica(self):
        tablas = {"ruta": TABLAS["ruta"]}
        wb = load_workbook(ReporteService(_esquema(tablas)).documentacion_excel())
        self.assertEqual(wb["ERD_TEXTUAL"]["A3"].value, "No se detectaron relaciones (FK).")

    def test_consulta_al_esquema_cada_tabla_una_vez(self):
        esquema = _esquema(TABLAS)
        ReporteService(esquema).documentacion_excel()
        self.assertEqual(esquema.ddl.call_count, 2)
        esquema.filas.assert_any_call("ruta")
        esquema.contar.assert_any_call("bus")
        esquema.claves_foraneas.assert_any_call("bus")


class TablaSinPkTest(unittest.TestCase):
    def setUp(self):
        tablas = {"log": {"ddl": "CREATE TABLE `log` (msg text)", "columnas": [_col("msg")],
                          "filas": [{"msg": "hola"}]}}
        self.wb = load_workbook(ReporteService(_esquema(tablas)).documentacion_excel())

    def tearDown(self):
        self.wb = None

    def test_indice_marca_estado_sin_pk_y_pk_vacia(self):
        fila = [c.value for c in self.wb["Índice"][2]]
        self.assertEqual((fila[4], fila[5], fila[6]), ("—", "—", "⚠ Sin PK"))

    def test_fila_sin_pk_se_resalta_con_el_relleno_de_alerta(self):
        for celda in self.wb["Índice"][2][:8]:
            self.assertTrue(celda.fill.fgColor.rgb.endswith("FFD7D7"))

    def test_hoja_de_tabla_muestra_guion_en_pk_y_fk(self):
        hoja = self.wb["log"]
        self.assertEqual((hoja["B6"].value, hoja["B7"].value), ("—", "—"))


class SinTablasTest(unittest.TestCase):
    def test_sin_tablas_solo_genera_las_hojas_fijas(self):
        wb = load_workbook(ReporteService(_esquema({})).documentacion_excel())
        self.assertEqual(wb.sheetnames, ["PORTADA", "Índice", "CREATE", "ERD_TEXTUAL"])
        self.assertEqual(wb["Índice"].max_row, 1)


if __name__ == "__main__":
    unittest.main()
