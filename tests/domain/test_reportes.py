import unittest
from datetime import datetime
from unittest.mock import Mock, patch

from openpyxl import Workbook

from app.services import reporte_service
from app.services.reporte_service import (
    BORDE,
    CENTRADO,
    ENCABEZADO,
    RELLENO_ENCABEZADO,
    ReporteService,
    _ajustar_columnas,
    _encabezado,
)


class NombreArchivoTest(unittest.TestCase):
    def setUp(self):
        self.esquema = Mock()
        self.esquema.nombre_bd.return_value = "stelman_db"
        self.servicio = ReporteService(self.esquema)

    def tearDown(self):
        self.esquema = self.servicio = None

    def test_nombre_sigue_el_formato_documentacion_bd_nombre_fecha(self):
        with patch.object(reporte_service, "datetime") as falso:
            falso.now.return_value = datetime(2026, 10, 24, 15, 30, 45)
            nombre = self.servicio.nombre_archivo()
        self.assertEqual(nombre, "DOCUMENTACION_BD_stelman_db_20261024_1530.xlsx")

    def test_nombre_real_cumple_el_patron_de_fecha_y_hora(self):
        nombre = self.servicio.nombre_archivo()
        self.assertRegex(nombre, r"^DOCUMENTACION_BD_stelman_db_\d{8}_\d{4}\.xlsx$")

    def test_nombre_usa_el_nombre_de_bd_que_entrega_el_esquema(self):
        self.esquema.nombre_bd.return_value = "otra_bd"
        self.assertIn("_otra_bd_", self.servicio.nombre_archivo())

    def test_nombre_termina_en_xlsx(self):
        self.assertTrue(self.servicio.nombre_archivo().endswith(".xlsx"))

    def test_nombre_consulta_el_nombre_de_bd_una_sola_vez(self):
        self.servicio.nombre_archivo()
        self.esquema.nombre_bd.assert_called_once_with()


class AjustarColumnasTest(unittest.TestCase):
    def setUp(self):
        self.hoja = Workbook().active

    def tearDown(self):
        self.hoja = None

    def _ancho(self, letra="A"):
        return self.hoja.column_dimensions[letra].width

    def test_texto_corto_usa_el_ancho_minimo_de_10(self):
        self.hoja["A1"] = "ab"
        _ajustar_columnas(self.hoja)
        self.assertEqual(self._ancho(), 10)

    def test_texto_de_8_caracteres_queda_justo_en_el_minimo(self):
        self.hoja["A1"] = "x" * 8
        _ajustar_columnas(self.hoja)
        self.assertEqual(self._ancho(), 10)

    def test_texto_de_9_caracteres_supera_el_minimo(self):
        self.hoja["A1"] = "x" * 9
        _ajustar_columnas(self.hoja)
        self.assertEqual(self._ancho(), 11)

    def test_texto_mediano_ocupa_largo_mas_dos(self):
        self.hoja["A1"] = "x" * 20
        _ajustar_columnas(self.hoja)
        self.assertEqual(self._ancho(), 22)

    def test_texto_muy_largo_se_limita_al_maximo_de_60(self):
        self.hoja["A1"] = "x" * 100
        _ajustar_columnas(self.hoja)
        self.assertEqual(self._ancho(), 60)

    def test_texto_de_58_caracteres_llega_justo_al_maximo(self):
        self.hoja["A1"] = "x" * 58
        _ajustar_columnas(self.hoja)
        self.assertEqual(self._ancho(), 60)

    def test_usa_la_celda_mas_larga_de_la_columna(self):
        self.hoja["A1"] = "corto"
        self.hoja["A2"] = "x" * 30
        self.hoja["A3"] = "medio"
        _ajustar_columnas(self.hoja)
        self.assertEqual(self._ancho(), 32)

    def test_convierte_numeros_a_texto_para_medir_el_largo(self):
        self.hoja["A1"] = 123456789012345
        _ajustar_columnas(self.hoja)
        self.assertEqual(self._ancho(), 17)

    def test_columna_con_celdas_vacias_usa_el_minimo(self):
        self.hoja.cell(row=1, column=1, value=None)
        _ajustar_columnas(self.hoja)
        self.assertEqual(self._ancho(), 10)

    def test_cada_columna_se_ajusta_por_separado(self):
        self.hoja["A1"] = "ab"
        self.hoja["B1"] = "x" * 25
        self.hoja["C1"] = "x" * 200
        _ajustar_columnas(self.hoja)
        self.assertEqual((self._ancho("A"), self._ancho("B"), self._ancho("C")), (10, 27, 60))


class EncabezadoTest(unittest.TestCase):
    def setUp(self):
        self.hoja = Workbook().active

    def tearDown(self):
        self.hoja = None

    def test_escribe_los_valores_en_columnas_consecutivas_desde_la_a(self):
        _encabezado(self.hoja, 1, ["#", "Tabla", "Registros"])
        valores = [self.hoja.cell(row=1, column=c).value for c in (1, 2, 3)]
        self.assertEqual(valores, ["#", "Tabla", "Registros"])

    def test_escribe_en_la_fila_indicada(self):
        _encabezado(self.hoja, 9, ["id"])
        self.assertEqual(self.hoja.cell(row=9, column=1).value, "id")
        self.assertIsNone(self.hoja.cell(row=1, column=1).value)

    def test_aplica_fuente_en_negrita(self):
        _encabezado(self.hoja, 1, ["Tabla"])
        self.assertTrue(self.hoja["A1"].font.bold)
        self.assertEqual(self.hoja["A1"].font, ENCABEZADO)

    def test_aplica_relleno_verde_de_encabezado(self):
        _encabezado(self.hoja, 1, ["Tabla"])
        celda = self.hoja["A1"]
        self.assertEqual(celda.fill, RELLENO_ENCABEZADO)
        self.assertEqual(celda.fill.fill_type, "solid")
        self.assertTrue(celda.fill.fgColor.rgb.endswith("C6E0B4"))

    def test_aplica_borde_fino_en_los_cuatro_lados(self):
        _encabezado(self.hoja, 1, ["Tabla"])
        borde = self.hoja["A1"].border
        self.assertEqual(borde, BORDE)
        for lado in (borde.left, borde.right, borde.top, borde.bottom):
            self.assertEqual(lado.style, "thin")

    def test_aplica_alineacion_centrada(self):
        _encabezado(self.hoja, 1, ["Tabla"])
        self.assertEqual(self.hoja["A1"].alignment, CENTRADO)
        self.assertEqual(self.hoja["A1"].alignment.horizontal, "center")
        self.assertEqual(self.hoja["A1"].alignment.vertical, "center")

    def test_todas_las_celdas_del_encabezado_tienen_el_mismo_estilo(self):
        _encabezado(self.hoja, 1, ["a", "b", "c"])
        for columna in (1, 2, 3):
            celda = self.hoja.cell(row=1, column=columna)
            self.assertEqual(
                (celda.font, celda.fill, celda.border, celda.alignment),
                (ENCABEZADO, RELLENO_ENCABEZADO, BORDE, CENTRADO),
            )

    def test_lista_vacia_no_escribe_celdas(self):
        _encabezado(self.hoja, 1, [])
        self.assertIsNone(self.hoja["A1"].value)


if __name__ == "__main__":
    unittest.main()
