"""Humo de punta a punta contra una BD real (MariaDB/MySQL con db/01 y db/02 cargados).

    RUN_INTEGRATION=1 pytest tests/integration

Cada prueba limpia lo que crea, así que se puede correr sobre la BD de desarrollo.
"""
import io
import os

import pymysql
import pytest
from openpyxl import load_workbook

from app import create_app
from app.config import Config

pytestmark = [
    pytest.mark.integration,
    pytest.mark.skipif(os.environ.get("RUN_INTEGRATION") != "1", reason="define RUN_INTEGRATION=1"),
]


@pytest.fixture(scope="module")
def app():
    class Testing(Config):
        TESTING = True
    return create_app(Testing)


@pytest.fixture()
def client(app):
    return app.test_client()


@pytest.fixture()
def sql(app):
    conn = pymysql.connect(host=app.config["DB_HOST"], port=app.config["DB_PORT"],
                           user=app.config["DB_USER"], password=app.config["DB_PASSWORD"],
                           database=app.config["DB_NAME"], charset="utf8mb4",
                           cursorclass=pymysql.cursors.DictCursor, autocommit=True)
    yield conn
    conn.close()


def consulta(sql, query, params=()):
    with sql.cursor() as cur:
        cur.execute(query, params or None)  # sin params, pymysql no debe interpretar los % del LIKE
        return cur.fetchall()


@pytest.fixture()
def sesion(client):
    r = client.post("/login", data={"usuario": "aperez_sup", "password": "stelman123"})
    assert r.status_code == 302 and r.headers["Location"].endswith("/panel")
    return client


# ---------------- público ----------------
def test_paginas_publicas_leen_rutas_de_la_bd(client):
    assert client.get("/").status_code == 200
    assert client.get("/contacto").status_code == 200
    assert b"EX-101" in client.get("/rutas").data
    detalle = client.get("/ruta/EX-101")
    assert detalle.status_code == 200 and b"Estacion Central" in detalle.data
    assert client.get("/ruta/NO-EXISTE").status_code == 404


def test_rutas_privadas_exigen_login(client):
    for url in ("/panel", "/panel/buses", "/panel/caja", "/exportar_excel", "/panel/personal/choferes"):
        r = client.get(url)
        assert r.status_code == 302 and "/login" in r.headers["Location"], url


# ---------------- login ----------------
def test_login_rechaza_clave_incorrecta(client):
    r = client.post("/login", data={"usuario": "aperez_sup", "password": "mala"}, follow_redirects=True)
    assert "incorrectos".encode() in r.data


def test_dashboard_muestra_ruta_del_supervisor(sesion):
    r = sesion.get("/panel")
    assert r.status_code == 200 and b"EX-101" in r.data


# ---------------- buses ----------------
def test_ciclo_de_buses(sesion, sql):
    try:
        r = sesion.post("/panel/buses", data={"placa": "tst-0001", "anio": "2024", "id_modelo_bus": "1"},
                        follow_redirects=True)
        assert b"Bus registrado" in r.data and b"TST-0001" in r.data
        dup = sesion.post("/panel/buses", data={"placa": "TST-0001", "anio": "2024", "id_modelo_bus": "1"},
                          follow_redirects=True)
        assert "ya está registrada".encode() in dup.data
        id_bus = consulta(sql, "SELECT id_bus FROM bus WHERE placa='TST-0001'")[0]["id_bus"]
        assert f"id_editar={id_bus}".encode() in sesion.get("/panel/buses").data
        assert b"Editar Bus: TST-0001" in sesion.get(f"/panel/buses?id_editar={id_bus}").data
        ed = sesion.post("/panel/buses", data={"id_bus_editado": id_bus, "placa": "TST-0002",
                                               "anio": "2025", "id_modelo_bus": "2"}, follow_redirects=True)
        assert b"Bus editado" in ed.data
        mal = sesion.post("/panel/buses", data={"placa": "TST-0003", "anio": "1800", "id_modelo_bus": "1"},
                          follow_redirects=True)
        assert b"El a\xc3\xb1o debe estar entre" in mal.data
        des = sesion.get(f"/panel/buses/eliminar/{id_bus}", follow_redirects=True)
        assert b"desasignado" in des.data and b"TST-0002" not in des.data
    finally:
        consulta(sql, "DELETE FROM bus_ruta WHERE id_bus IN (SELECT id_bus FROM bus WHERE placa LIKE 'TST-%')")
        consulta(sql, "DELETE FROM bus WHERE placa LIKE 'TST-%'")


def test_no_puede_tocar_bus_de_otra_ruta(sesion):
    r = sesion.get("/panel/buses/eliminar/2", follow_redirects=True)  # bus 2 no está en la ruta 1
    assert b"Acceso denegado" in r.data


# ---------------- choferes ----------------
def test_ciclo_de_choferes(sesion, sql):
    form = {"nombre": "Luis", "apellido": "Prueba", "dni": "TST-DNI-1", "fecha_nacimiento": "1990-01-01",
            "correo": "", "telefono": "999", "sueldo": "1800", "fecha_ingreso": "2026-01-10",
            "nro_licencia": "TST-LIC-1", "fecha_primera_licencia": "2015-03-01", "id_tipo_licencia": "1"}
    try:
        assert b"Carlos" in sesion.get("/panel/personal/choferes").data
        r = sesion.post("/registrar_personal_chofer", data=form, follow_redirects=True)
        assert b"Chofer registrado" in r.data and b"Prueba" in r.data
        dup = sesion.post("/registrar_personal_chofer", data=form, follow_redirects=True)
        assert b"Ya existe" in dup.data
        id_emp = consulta(sql, "SELECT id_empleado FROM empleado WHERE nro_identificacion='TST-DNI-1'")[0]["id_empleado"]
        assert b"Editar Chofer" in sesion.get(f"/panel/personal/choferes?editar={id_emp}").data
        upd = sesion.post("/actualizar_personal_chofer", data={**form, "id_empleado": id_emp, "sueldo": "2100"},
                          follow_redirects=True)
        assert b"actualizados" in upd.data
        assert consulta(sql, "SELECT sueldo FROM empleado WHERE id_empleado=%s", (id_emp,))[0]["sueldo"] == 2100
        eli = sesion.get(f"/panel/personal/choferes?eliminar={id_emp}", follow_redirects=True)
        assert b"Registro eliminado" in eli.data
        assert not consulta(sql, "SELECT 1 FROM empleado WHERE id_empleado=%s", (id_emp,))
    finally:
        consulta(sql, "DELETE FROM chofer WHERE nro_licencia='TST-LIC-1'")
        consulta(sql, "DELETE FROM empleado WHERE nro_identificacion='TST-DNI-1'")


def test_no_elimina_chofer_con_historial(sesion, sql):
    r = sesion.get("/panel/personal/choferes?eliminar=1", follow_redirects=True)  # Carlos tiene caja/asignaciones
    assert b"No se puede eliminar" in r.data
    assert consulta(sql, "SELECT 1 FROM empleado WHERE id_empleado=1")


# ---------------- incidencias ----------------
def test_ciclo_incidencia_disciplinaria(sesion, sql):
    form = {"fecha": "2026-09-01", "descripcion": "TST impuntualidad", "estado": "Pendiente",
            "id_bus": "1", "id_tipo_disciplinaria": "2", "sancion": "Amonestación"}
    try:
        page = sesion.get("/panel/incidencias/disciplinarias?nuevo_incidencia=1")
        assert page.status_code == 200 and b"Impuntualidad grave" in page.data
        r = sesion.post("/registrar_incidencia_disciplinaria", data=form, follow_redirects=True)
        assert b"registrada" in r.data and b"TST impuntualidad" in r.data
        id_inc = consulta(sql, "SELECT id_incidencia FROM incidencia WHERE descripcion='TST impuntualidad'")[0]["id_incidencia"]
        assert b'value="2026-09-01"' in sesion.get(f"/panel/incidencias/disciplinarias?editar={id_inc}").data
        up = sesion.post("/actualizar_incidencia_disciplinaria",
                         data={**form, "id_incidencia": id_inc, "estado": "Resuelta"}, follow_redirects=True)
        assert b"actualizada" in up.data
        futuro = sesion.post("/registrar_incidencia_disciplinaria", data={**form, "fecha": "2999-01-01"},
                             follow_redirects=True)
        assert b"fuera del rango" in futuro.data
        assert b"eliminada" in sesion.get(f"/eliminar_incidencia_disciplinaria/{id_inc}", follow_redirects=True).data
    finally:
        consulta(sql, "DELETE FROM incidencia_disciplinaria WHERE id_incidencia IN (SELECT id_incidencia FROM incidencia WHERE descripcion LIKE 'TST%')")
        consulta(sql, "DELETE FROM incidencia WHERE descripcion LIKE 'TST%'")


def test_ciclo_incidencia_operativa(sesion, sql):
    form = {"fecha": "2026-09-02", "descripcion": "TST falla de frenos", "estado": "En revision",
            "id_bus": "1", "gravedad": "Grave", "costo": "250.50", "requiere_seguro": "on"}
    try:
        assert b"Choque leve" in sesion.get("/panel/incidencias/operativas").data  # incidencia del seed
        r = sesion.post("/registrar_incidencia_operativa", data=form, follow_redirects=True)
        assert b"registrada" in r.data and b"TST falla de frenos" in r.data
        id_inc = consulta(sql, "SELECT id_incidencia FROM incidencia WHERE descripcion='TST falla de frenos'")[0]["id_incidencia"]
        up = sesion.post("/actualizar_incidencia_operativa",
                         data={**form, "id_incidencia": id_inc, "gravedad": "Leve"}, follow_redirects=True)
        assert b"actualizada" in up.data
        assert consulta(sql, "SELECT gravedad FROM incidencia_operativa WHERE id_incidencia=%s", (id_inc,))[0]["gravedad"] == "Leve"
        assert b"eliminada" in sesion.get(f"/eliminar_incidencia_operativa/{id_inc}", follow_redirects=True).data
    finally:
        consulta(sql, "DELETE FROM incidencia_operativa WHERE id_incidencia IN (SELECT id_incidencia FROM incidencia WHERE descripcion LIKE 'TST%')")
        consulta(sql, "DELETE FROM incidencia WHERE descripcion LIKE 'TST%'")


# ---------------- caja ----------------
def test_ciclo_caja(sesion, sql):
    form = {"fecha": "2026-09-03", "monto": "410.75", "observacion": "TST caja", "id_empleado": "1", "id_bus": "1"}
    try:
        assert b"350.50" in sesion.get("/panel/caja").data
        r = sesion.post("/registrar_caja", data=form, follow_redirects=True)
        assert b"Recaudaci\xc3\xb3n registrada" in r.data and b"410.75" in r.data
        id_caja = consulta(sql, "SELECT id_caja FROM caja WHERE observacion='TST caja'")[0]["id_caja"]
        assert b'value="2026-09-03"' in sesion.get(f"/editar_caja/{id_caja}").data
        up = sesion.post("/actualizar_caja", data={**form, "id_caja": id_caja, "monto": "500"}, follow_redirects=True)
        assert b"actualizada" in up.data
        cero = sesion.post("/registrar_caja", data={**form, "monto": "0"}, follow_redirects=True)
        assert b"El monto debe ser mayor" in cero.data
        assert b"eliminada" in sesion.get(f"/eliminar_caja/{id_caja}", follow_redirects=True).data
    finally:
        consulta(sql, "DELETE FROM caja WHERE observacion='TST caja'")


# ---------------- reportes ----------------
def test_exporta_excel_con_todas_las_tablas(sesion, sql):
    r = sesion.get("/exportar_excel")
    assert r.status_code == 200 and "spreadsheetml" in r.mimetype
    hojas = load_workbook(io.BytesIO(r.data)).sheetnames
    tablas = {f["TABLE_NAME"] for f in consulta(sql, "SELECT TABLE_NAME FROM information_schema.TABLES WHERE TABLE_SCHEMA=DATABASE()")}
    assert {"PORTADA", "Índice", "CREATE", "ERD_TEXTUAL"} <= set(hojas)
    assert tablas <= set(hojas)
