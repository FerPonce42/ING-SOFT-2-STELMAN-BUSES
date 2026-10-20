from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITULO = Font(size=20, bold=True)
SUBTITULO = Font(size=12, italic=True)
ENCABEZADO = Font(bold=True)
RELLENO_PK = PatternFill("solid", fgColor="FFD966")
RELLENO_FK = PatternFill("solid", fgColor="C9DAF8")
RELLENO_DDL = PatternFill("solid", fgColor="FFF2CC")
RELLENO_ENCABEZADO = PatternFill("solid", fgColor="C6E0B4")
RELLENO_ALERTA = PatternFill("solid", fgColor="FFD7D7")
_LADO = Side(style="thin")
BORDE = Border(left=_LADO, right=_LADO, top=_LADO, bottom=_LADO)
CENTRADO = Alignment(horizontal="center", vertical="center")


def _ajustar_columnas(hoja):
    for columna in hoja.columns:
        largo = max((len(str(c.value)) for c in columna if c.value is not None), default=0)
        hoja.column_dimensions[get_column_letter(columna[0].column)].width = min(max(largo + 2, 10), 60)


def _encabezado(hoja, fila, valores):
    for i, valor in enumerate(valores, start=1):
        celda = hoja.cell(row=fila, column=i, value=valor)
        celda.font, celda.fill, celda.border, celda.alignment = ENCABEZADO, RELLENO_ENCABEZADO, BORDE, CENTRADO


class ReporteService:
    def __init__(self, esquema):
        self.esquema = esquema

    def nombre_archivo(self):
        return f"DOCUMENTACION_BD_{self.esquema.nombre_bd()}_{datetime.now():%Y%m%d_%H%M}.xlsx"

    def documentacion_excel(self):
        esquema = self.esquema
        wb = Workbook()
        wb.remove(wb.active)

        portada = wb.create_sheet("PORTADA")
        portada.merge_cells("A1:F1")
        portada["A1"] = "DOCUMENTACIÓN TÉCNICA - BD STELMAN BUSES"
        portada["A1"].font = TITULO
        portada["A3"] = f"Base de datos: {esquema.nombre_bd()}"
        portada["A3"].font = SUBTITULO
        portada["A4"] = f"Generado el: {datetime.now():%Y-%m-%d %H:%M:%S}"
        portada["A5"] = "Generado por: Sistema Stelman"
        portada["A7"] = "Contenido:"
        portada["A8"] = "- Índice con links a cada hoja"
        portada["A9"] = "- DDL completo por tabla"
        portada["A10"] = "- Hojas por tabla con datos, PK y FK resaltadas"
        _ajustar_columnas(portada)

        indice = wb.create_sheet("Índice")
        _encabezado(indice, 1, ["#", "Tabla", "Atributos", "Registros", "PK", "FK", "Estado", "Link hoja"])
        hoja_ddl = wb.create_sheet("CREATE")
        _encabezado(hoja_ddl, 1, ["Tabla", "CREATE TABLE (DDL)"])

        relaciones = []
        for numero, tabla in enumerate(esquema.tablas(), start=1):
            ddl = esquema.ddl(tabla)
            columnas = esquema.columnas(tabla)
            fks = esquema.claves_foraneas(tabla)
            pk = [c["COLUMN_NAME"] for c in columnas if c["COLUMN_KEY"] == "PRI"]
            columnas_fk = {f["COLUMN_NAME"] for f in fks}
            fk_texto = ", ".join(
                f"{f['COLUMN_NAME']} → {f['REFERENCED_TABLE_NAME']}.{f['REFERENCED_COLUMN_NAME']}" for f in fks
            ) or "—"
            estado = "OK" if pk else "⚠ Sin PK"

            indice.append([numero, tabla, len(columnas), esquema.contar(tabla),
                           ", ".join(pk) or "—", fk_texto, estado,
                           f'=HYPERLINK("#{tabla}!A1","Abrir")'])
            hoja_ddl.append([tabla, ddl])
            relaciones += [(tabla, f["COLUMN_NAME"], f["REFERENCED_TABLE_NAME"], f["REFERENCED_COLUMN_NAME"]) for f in fks]

            ws = wb.create_sheet(tabla[:31])
            ws["A1"] = f"CREATE TABLE: {tabla}"
            ws["A1"].font, ws["A1"].fill = ENCABEZADO, RELLENO_DDL
            ws.merge_cells("A1:F1")
            ws["A2"] = ddl
            ws["A2"].alignment, ws["A2"].fill = Alignment(wrap_text=True), RELLENO_DDL
            ws.merge_cells("A2:F4")
            ws.row_dimensions[2].height = 70
            ws["A6"], ws["B6"] = "PK:", ", ".join(pk) or "—"
            ws["A7"], ws["B7"] = "FK:", fk_texto

            nombres = [c["COLUMN_NAME"] for c in columnas]
            _encabezado(ws, 9, nombres)
            filas = esquema.filas(tabla)
            for r, fila in enumerate(filas, start=10):
                for c, nombre in enumerate(nombres, start=1):
                    celda = ws.cell(row=r, column=c, value=fila.get(nombre))
                    celda.border = BORDE
                    if nombre in pk:
                        celda.fill = RELLENO_PK
                    elif nombre in columnas_fk:
                        celda.fill = RELLENO_FK
            if filas:
                ws.auto_filter.ref = f"A9:{get_column_letter(len(nombres))}{9 + len(filas)}"
            ws.freeze_panes = "A10"
            _ajustar_columnas(ws)

        erd = wb.create_sheet("ERD_TEXTUAL")
        erd["A1"] = "Mini ERD (texto)"
        erd["A1"].font = ENCABEZADO
        if relaciones:
            for r, (t, c, td, cd) in enumerate(relaciones, start=3):
                erd.cell(row=r, column=1, value=f"{t}.{c}  →  {td}.{cd}")
        else:
            erd["A3"] = "No se detectaron relaciones (FK)."
        _ajustar_columnas(erd)

        _ajustar_columnas(indice)
        _ajustar_columnas(hoja_ddl)
        for fila in indice.iter_rows(min_row=2, max_col=8):
            if "Sin PK" in str(fila[6].value):
                for celda in fila:
                    celda.fill = RELLENO_ALERTA

        salida = BytesIO()
        wb.save(salida)
        salida.seek(0)
        return salida
